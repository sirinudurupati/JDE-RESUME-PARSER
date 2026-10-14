import io

from conftest import make_docx, make_pdf
from talentmatch.ats import ats_report
from talentmatch.extract import extract_text, inspect_file
from talentmatch.resume_parser import parse_resume

GOOD = """Jane Roe
jane@example.com | +91 98765 43210
SUMMARY
Backend developer.
EXPERIENCE
Python Developer
Acme Technologies Pvt Ltd
Jan 2022 - Dec 2023
- Built APIs serving 10000 users and cut latency by 30%
- Wrote 120 unit tests with pytest and reduced bugs by 25%
- Migrated 4 services to Docker and saved 15% hosting cost
EDUCATION
B.Tech in Computer Science, ABC University 2018 - 2022
SKILLS
Python, SQL, FastAPI, Docker
"""


def report(text, data=None, ext=".pdf"):
    p = parse_resume(text)
    if data is not None:
        p["file_info"] = inspect_file(data, ext)
    return ats_report(p)


def test_clean_resume_scores_high_with_all_checks_explained():
    r = report(GOOD, make_pdf(GOOD))
    assert r["score"] >= 90 and r["band"] == "Good"
    for c in r["checks"]:
        assert c["detail"] and 0 <= c["points"] <= c["max"]
    assert sum(c["max"] for c in r["checks"]) == 100


def test_two_column_pdf_is_detected_and_penalised():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for i in range(40):
        c.drawString(40, y, f"Left column line {i} with several words here")
        c.drawString(330, y, f"Right column line {i} with several words here")
        y -= 16
    c.save()
    data = buf.getvalue()
    assert inspect_file(data, ".pdf")["multi_column"] is True
    single = make_pdf("\n".join(f"Single column line {i} with several words in it for testing" for i in range(40)))
    assert inspect_file(single, ".pdf")["multi_column"] is False
    r = report(GOOD, data)
    assert next(k for k in r["checks"] if k["key"] == "columns")["status"] == "fail"


def test_docx_table_and_pages_detected():
    import docx
    d = docx.Document()
    d.add_paragraph("Jane Roe")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Skills"
    buf = io.BytesIO()
    d.save(buf)
    info = inspect_file(buf.getvalue(), ".docx")
    assert info["tables"] == 1
    r = report(GOOD, buf.getvalue(), ".docx")
    assert next(k for k in r["checks"] if k["key"] == "tables")["status"] == "fail"


def test_image_only_text_fails_readable_text_check():
    r = ats_report(parse_resume("x"))
    assert next(k for k in r["checks"] if k["key"] == "text")["status"] == "fail"
    assert r["score"] < 40


def test_unknown_layout_is_not_silently_full_marks():
    r = ats_report(parse_resume(GOOD))                 # no file available
    lay = next(k for k in r["checks"] if k["key"] == "layout")
    assert lay["status"] == "unknown" and r["layout_assessed"] is False
    assert r["assessed_points"] == 80                  # 20 layout points left out of the denominator


def test_emoji_garbled_spaced_and_missing_sections_flagged():
    bad = GOOD.replace("jane@example.com", "✉ jane@example.com").replace("EXPERIENCE", "My Journey") \
        + "\nN o v e m b e r\n"
    r = report(bad, make_pdf(GOOD))
    keys = {c["key"]: c for c in r["checks"]}
    assert keys["chars"]["status"] == "warn" and keys["chars"]["points"] < 10
    assert keys["sections"]["points"] < 20


def test_scores_in_range_and_deterministic():
    a, b = report(GOOD, make_pdf(GOOD)), report(GOOD, make_pdf(GOOD))
    assert a == b and 0 <= a["score"] <= 100


def test_ats_in_candidate_result(client, jd_text):
    r = client.post("/api/jobs/upload", data={"file": (io.BytesIO(jd_text.encode()), "jd.txt")},
                    content_type="multipart/form-data").get_json()["job"]["id"]
    out = client.post(f"/api/jobs/{r}/resumes",
                      data={"files": [(io.BytesIO(make_pdf(GOOD)), "cv.pdf"), (io.BytesIO(make_docx(GOOD)), "cv.docx")]},
                      content_type="multipart/form-data").get_json()["results"]
    assert all(x["status"] == "ok" for x in out)
    cands = client.get(f"/api/jobs/{r}/candidates").get_json()["candidates"]
    for c in cands:
        ats = c["result"]["ats"]
        assert 0 <= ats["score"] <= 100 and ats["layout_assessed"] is True and ats["keywords"]["total"] > 0
    csv_text = client.get(f"/api/jobs/{r}/export.csv").get_data(as_text=True)
    assert "ATS score" in csv_text
