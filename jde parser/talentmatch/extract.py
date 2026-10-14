"""Text extraction + upload validation for PDF / DOCX / TXT."""
import io
import unicodedata
import zipfile


class ExtractionError(Exception):
    """Raised with a user-presentable message."""


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractionError("PDF is password protected.")
            except ExtractionError:
                raise
            except Exception:
                raise ExtractionError("PDF is password protected.")
        parts = [(p.extract_text() or "") for p in reader.pages]
    except ExtractionError:
        raise
    except Exception:
        raise ExtractionError("PDF is corrupted or unreadable.")
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    import docx
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise ExtractionError("DOCX is corrupted or unreadable.")
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in seen:
                    seen.append(t)
            if seen:
                lines.append(" | ".join(seen))
    return "\n".join(lines)


def _txt_text(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise ExtractionError("File is not a plain-text document.")
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ExtractionError("Text file encoding not supported.")


def sniff_ext(data: bytes, ext: str) -> str:
    """Validate that content matches the claimed extension (magic bytes)."""
    ext = ext.lower()
    if ext == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF"):
            raise ExtractionError("File is not a valid PDF.")
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise ExtractionError("File is not a valid DOCX.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise ExtractionError("File is not a valid DOCX.")
        except zipfile.BadZipFile:
            raise ExtractionError("DOCX is corrupted or unreadable.")
    return ext


def extract_text(data: bytes, ext: str, allowed: set) -> str:
    ext = ext.lower()
    if ext not in allowed:
        raise ExtractionError(
            f"Unsupported file type '{ext or 'unknown'}'. Allowed: "
            + ", ".join(sorted(allowed)).replace(".", "").upper())
    if not data:
        raise ExtractionError("File is empty.")
    sniff_ext(data, ext)
    text = {".pdf": _pdf_text, ".docx": _docx_text, ".txt": _txt_text}[ext](data)
    # NFKC folds styled letters (e.g. bold "Mandatory"), ligatures and odd spaces to plain text
    text = unicodedata.normalize("NFKC", text).replace(" ", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    if len(text.strip()) < 20:
        raise ExtractionError(
            "No readable text found (the file may be a scanned image or empty).")
    return text
