"""CSV / Excel candidate comparison reports."""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Rank", "Candidate", "Email", "Phone", "Highest education", "Graduation year",
           "Total experience (yrs, dated)", "Stated experience (months, undated)", "Certifications",
           "Overall match (0-100)", "ATS score (0-100)", "ATS issues", "Skills %",
           "Relevant experience %", "Education %", "Responsibilities %", "Certifications %",
           "Total experience (months)", "Relevant experience (months)", "Required experience (months)",
           "Missing experience (months)", "Matched mandatory",
           "Related-only mandatory", "Missing mandatory", "Matched preferred", "Missing preferred",
           "Needs manual review", "Review reasons", "Resume file"]


def _cat(res, key):
    c = res["categories"][key]
    return c["score"] if c["applicable"] else "n/a"


def _grad(p):
    if not p or not p.get("graduation_year"):
        return ""
    return f"{p['graduation_year']}" + (" (expected)" if p.get("graduation_expected") else "")


def rows_for(candidates):
    ranked = sorted([c for c in candidates if c.get("result")],
                    key=lambda c: (-c["overall"], (c["name"] or "").lower()))
    rows = []
    for i, c in enumerate(ranked, 1):
        r = c["result"]
        exp = r["experience"]
        rows.append([
            i, c["name"] or "(name not detected)", c["email"] or "", c["phone"] or "",
            (r.get("profile") or {}).get("highest_education") or "",
            _grad(r.get("profile")),
            (r.get("profile") or {}).get("experience_years", ""),
            (r.get("profile") or {}).get("stated_months", ""),
            "; ".join((r.get("profile") or {}).get("certifications", [])),
            r["overall"], (r.get("ats") or {}).get("score", ""), "; ".join((r.get("ats") or {}).get("issues", [])),
            _cat(r, "skills"), _cat(r, "experience"), _cat(r, "education"),
            _cat(r, "responsibilities"), _cat(r, "certifications"),
            exp.get("total_months", ""), exp.get("relevant_months", ""),
            exp.get("required_months") or "n/a",
            exp.get("missing_months") if exp.get("required_months") else "n/a",
            "; ".join(r["matched_mandatory"]), "; ".join(r["related_mandatory"]),
            "; ".join(r["missing_mandatory"]), "; ".join(r["matched_preferred"]),
            "; ".join(r["missing_preferred"]),
            "Yes" if r["needs_review"] else "No", "; ".join(r["review_reasons"]), c["filename"]])
    return rows


def _safe(v):
    """Neutralise spreadsheet formula injection from resume-derived text."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def to_csv(job, candidates) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Job", job["title"]])
    w.writerow([])
    w.writerow(HEADERS)
    for row in rows_for(candidates):
        w.writerow([_safe(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(job, candidates) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    ws.append(["Job", job["title"]])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(HEADERS)
    head = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows_for(candidates):
        ws.append([_safe(v) for v in row])
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(h) + 4))
    ws.freeze_panes = "C4"

    info = wb.create_sheet("Job & weights")
    req = job["requirements"]
    info.append(["Title", job["title"]])
    info.append(["Department", req.get("department", "")])
    info.append(["Minimum experience (yrs)", req.get("min_experience_years") or "not set"])
    info.append([])
    info.append(["Category", "Configured weight"])
    for k, v in job["weights"].items():
        info.append([k, v])
    info.append([])
    info.append(["Skill", "Level"])
    for s in req.get("skills", []):
        info.append([s["display"], s["level"]])
    info.append([])
    info.append(["Note", "Match scores estimate job alignment for human review; they are not "
                         "hiring decisions."])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 40
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
