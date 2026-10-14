"""Text extraction + upload validation for PDF / DOCX / TXT."""
import io
import logging
import unicodedata
import zipfile

logging.getLogger("pdfminer").setLevel(logging.ERROR)   # silence harmless font warnings


class ExtractionError(Exception):
    """Raised with a user-presentable message."""


def _pdf_text_plumber(data: bytes) -> str:
    """pdfplumber keeps lines, spacing and column order much better than pypdf on
    designed resumes (dates on the same line as the institution, spaced-out letters...)."""
    import pdfplumber
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if getattr(pdf, "doc", None) is not None and getattr(pdf.doc, "is_encrypted", False):
            raise ExtractionError("PDF is password protected.")
        for page in pdf.pages:
            pages.append(page.extract_text(x_tolerance=2, y_tolerance=3) or "")
    return "\n".join(pages)


def _pdf_text(data: bytes) -> str:
    try:
        text = _pdf_text_plumber(data)
        if len(text.strip()) >= 20:
            return text
    except ExtractionError:
        raise
    except Exception:
        pass                      # fall back to pypdf below
    return _pdf_text_pypdf(data)


def _pdf_text_pypdf(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractionError("PDF is password protected.")
            except ExtractionError:
                raise
            except Exception:
                raise ExtractionError("PDF is password protected.")
        parts = [(p.extract_text() or "") for p in reader.pages]
    except ExtractionError:
        raise
    except Exception:
        raise ExtractionError("PDF is corrupted or unreadable.")
    return "\n".join(parts)


def _docx_text(data: bytes) -> str:
    import docx
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception:
        raise ExtractionError("DOCX is corrupted or unreadable.")
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in seen:
                    seen.append(t)
            if seen:
                lines.append(" | ".join(seen))
    return "\n".join(lines)


def _txt_text(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise ExtractionError("File is not a plain-text document.")
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ExtractionError("Text file encoding not supported.")


def sniff_ext(data: bytes, ext: str) -> str:
    """Validate that content matches the claimed extension (magic bytes)."""
    ext = ext.lower()
    if ext == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF"):
            raise ExtractionError("File is not a valid PDF.")
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise ExtractionError("File is not a valid DOCX.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise ExtractionError("File is not a valid DOCX.")
        except zipfile.BadZipFile:
            raise ExtractionError("DOCX is corrupted or unreadable.")
    return ext


def extract_text(data: bytes, ext: str, allowed: set) -> str:
    ext = ext.lower()
    if ext not in allowed:
        raise ExtractionError(
            f"Unsupported file type '{ext or 'unknown'}'. Allowed: "
            + ", ".join(sorted(allowed)).replace(".", "").upper())
    if not data:
        raise ExtractionError("File is empty.")
    sniff_ext(data, ext)
    text = {".pdf": _pdf_text, ".docx": _docx_text, ".txt": _txt_text}[ext](data)
    # NFKC folds styled letters (e.g. bold "Mandatory"), ligatures and odd spaces to plain text
    text = unicodedata.normalize("NFKC", text).replace(" ", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    if len(text.strip()) < 20:
        raise ExtractionError(
            "No readable text found (the file may be a scanned image or empty).")
    return text


# ---------------------------------------------------------------- file inspection (for ATS checks)
def _pdf_columns(page) -> bool:
    """True when the page looks multi-column: many text lines have a gap at the SAME x position with
    several words on both sides. (Right-aligned dates leave only 1-3 words on the right, so they pass.)"""
    words = page.extract_words(x_tolerance=2, y_tolerance=3) or []
    if len(words) < 40:
        return False
    width = float(page.width)
    lines = {}
    for w in words:
        lines.setdefault(round(w["top"] / 4), []).append(w)
    centres, total = [], 0
    for ws in lines.values():
        if len(ws) < 8:
            continue
        total += 1
        ws.sort(key=lambda w: w["x0"])
        best = max(range(1, len(ws)), key=lambda i: ws[i]["x0"] - ws[i - 1]["x1"])
        gap = ws[best]["x0"] - ws[best - 1]["x1"]
        left, right = best, len(ws) - best
        mid = (ws[best]["x0"] + ws[best - 1]["x1"]) / 2
        if gap > width * 0.03 and left >= 4 and right >= 4 and 0.25 * width < mid < 0.75 * width:
            centres.append(mid)
    if total < 6 or len(centres) < max(5, 0.4 * total):
        return False
    centres.sort()
    median = centres[len(centres) // 2]
    return sum(1 for c in centres if abs(c - median) < width * 0.04) >= max(5, 0.4 * total)


def inspect_file(data: bytes, ext: str) -> dict:
    """Layout facts that affect ATS parsing. Never raises; missing facts stay None."""
    import re
    info = {"ext": ext.lower(), "pages": None, "tables": None, "images": None, "multi_column": None,
            "textboxes": None, "header_footer_contact": None, "size_kb": round(len(data) / 1024, 1)}
    try:
        if ext.lower() == ".pdf":
            import pdfplumber
            tables = images = 0
            cols = False
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                info["pages"] = len(pdf.pages)
                for page in pdf.pages[:4]:
                    tables += len(page.find_tables() or [])
                    images += len(page.images or [])
                    cols = cols or _pdf_columns(page)
            info.update(tables=tables, images=images, multi_column=cols, textboxes=0,
                        header_footer_contact=False)
        elif ext.lower() == ".docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                body = z.read("word/document.xml").decode("utf8", "ignore")
                app = z.read("docProps/app.xml").decode("utf8", "ignore") if "docProps/app.xml" in z.namelist() else ""
            m = re.search(r"<Pages>(\d+)</Pages>", app)
            hf = False
            for sec in d.sections:
                for part in (sec.header, sec.footer):
                    t = " ".join(p.text for p in part.paragraphs)
                    if re.search(r"@|\+?\d[\d\s().-]{8,}\d", t):
                        hf = True
            info.update(pages=int(m.group(1)) if m else None, tables=len(d.tables),
                        images=len(d.inline_shapes), multi_column=bool(re.search(r'<w:cols [^>]*w:num="[2-9]"', body)),
                        textboxes=body.count("<w:txbxContent"), header_footer_contact=hf)
    except Exception:
        pass
    return info
